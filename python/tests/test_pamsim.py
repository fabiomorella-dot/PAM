import json
import os
import shutil
import sys

import numpy as np
import pytest
from scipy.special import erf

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
from pamsim import SimSettings, Species, simulate, cross_correlation, correlate_blocks, load_sim  # noqa: E402
from pamsim.correlation import do_ccf, time_axis  # noqa: E402

HAVE_COMPILERS = shutil.which("gcc") is not None and shutil.which("g++") is not None


def small(**kw):
    kw.setdefault("box_size", [2000, 2000, 3000])
    kw.setdefault("sim_time", 0.4)
    kw.setdefault("frames", 2)
    kw.setdefault("seed", 1)
    return SimSettings(**kw)


def box_rate(brightness_hz, n, box, wr=200.0, wz=1000.0):
    """Expected count rate for Gaussian excitation in a box (focus in the centre)."""
    w = wr * np.sqrt(np.pi / 2) * erf(np.sqrt(2) * box[0] / 2 / wr)
    h = wr * np.sqrt(np.pi / 2) * erf(np.sqrt(2) * box[1] / 2 / wr)
    z = wz * np.sqrt(np.pi / 2) * erf(np.sqrt(2) * box[2] / 2 / wz)
    return brightness_hz * n / np.prod(box) * w * h * z


# ---------------------------------------------------------------- RNG / API
def test_reproducible_and_thread_independent():
    s = small(species=[Species(D=40, N=6, ExP=np.eye(4) * 0.3)])
    s.n_threads = 1
    a = simulate(s)
    s.n_threads = 4
    b = simulate(s)
    assert a.seed == b.seed == 1
    for c in range(4):
        np.testing.assert_array_equal(a.macrotimes[c], b.macrotimes[c])
        np.testing.assert_array_equal(a.microtimes[c], b.microtimes[c])
    s.seed = 2
    c = simulate(s)
    assert c.macrotimes[0].size != a.macrotimes[0].size or not np.array_equal(c.macrotimes[0], a.macrotimes[0])


def test_random_seed_is_recorded():
    s = small(seed=None, sim_time=0.05, species=[Species(D=40, N=2, ExP=np.eye(4) * 0.3)])
    r = simulate(s)
    s.seed = r.seed
    r2 = simulate(s)
    np.testing.assert_array_equal(r.macrotimes[0], r2.macrotimes[0])
    assert r.header["Info"]["General"]["SeedString"] == str(r.seed)


def test_output_layout():
    exp_ = np.eye(4) * 0.3
    cross = np.eye(4); cross[1, 0] = 0.2
    s = small(species=[Species(D=40, N=4, color=2, fret=2, ExP=exp_, Cross=cross)])
    r = simulate(s)
    assert r.n_channels == 4
    for c in range(4):
        assert np.all(np.diff(r.macrotimes[c]) >= 0)
        assert np.all((r.labels[c] & 3) == c)
        assert r.macrotimes[c].size == r.microtimes[c].size == r.labels[c].size
    total = s.frames * r.header["FrameTime"]
    assert r.macrotimes[0].max() < total
    assert r.header["MI_Bins"] == 8192 and r.header["Freq"] == 1e6


# ---------------------------------------------------------------- physics
@pytest.mark.parametrize("ff", [True, False])
def test_count_rate_matches_theory(ff):
    rates = []
    for k in range(16):
        s = small(seed=100 + k, fast_forward=ff, species=[Species(D=40, N=40, ExP=np.eye(4) * 0.3)])
        rates.append(simulate(s).macrotimes[0].size / s.sim_time)
    rates = np.array(rates)
    expected = box_rate(300e3, 40, np.array([2000, 2000, 3000.0]))
    se = rates.std(ddof=1) / np.sqrt(rates.size)
    assert abs(rates.mean() - expected) < 4.5 * se, (rates.mean(), expected, se)


def test_fcs_amplitude_and_diffusion_time():
    from scipy.optimize import curve_fit
    D, wr, wz, N = 50.0, 200.0, 1000.0, 40
    box = np.array([4000.0, 4000.0, 6000.0])
    s = SimSettings(box_size=box, sim_time=20, frames=1, seed=3,
                    species=[Species(D=D, N=N, brightness=[100] * 4)])
    r = simulate(s)
    lags, g, sem, _ = correlate_blocks(r.macrotimes[0], r.macrotimes[0], 20e6, n_blocks=10)
    tau = lags * 1e-6
    model = lambda t, n, td, g0: g0 + 1 / n / (1 + t / td) / np.sqrt(1 + t / (25 * td))  # noqa: E731
    ok = (tau > 2e-6) & (tau < 0.05) & (sem > 0)
    (n, td, _), _ = curve_fit(model, tau[ok], g[ok], p0=[0.1, 2e-4, 0], sigma=sem[ok])
    n_theory = N / np.prod(box * 1e-3) * np.pi ** 1.5 * (wr * 1e-3) ** 2 * wz * 1e-3
    td_theory = (wr * 1e-3) ** 2 / (4 * D)
    assert abs(n / n_theory - 1) < 0.08
    assert abs(td / td_theory - 1) < 0.12


def test_lifetime_and_pie_offset():
    s = small(species=[Species(D=40, N=10, color=2, ExP=np.eye(4) * 0.3, LT=[4, 2, 1, 1])])
    r = simulate(s)
    lt_bins = np.array([4, 2]) / s.mi_range * 8192
    mi0 = r.microtimes[0].astype(float) - 1                 # donor, laser 1: offset 1
    mi1 = r.microtimes[1].astype(float) - (16384 + 1)       # acceptor, laser 2: offset 16384+1
    for mi, lt in ((mi0, lt_bins[0]), (mi1, lt_bins[1])):
        p = 1 / lt                                          # geometric distribution, mean (1-p)/p
        assert abs(mi.mean() - (1 - p) / p) < 5 * np.sqrt((1 - p) / p ** 2 / mi.size)


def test_fret_efficiency():
    R = np.zeros((4, 4)); R[1, 0] = 63.0     # R = R0  ->  E = 0.5
    exp_ = np.zeros((4, 4)); exp_[0, 0] = 0.3
    s = small(species=[Species(D=40, N=10, color=2, fret=2, ExP=exp_, R=R)])
    r = simulate(s)
    nd, na = r.macrotimes[0].size, r.macrotimes[1].size
    e = na / (nd + na)
    assert abs(e - 0.5) < 5 * np.sqrt(0.25 / (nd + na))


def test_background_noise_rate():
    s = small(sim_time=2.0, use_noise=True, noise=[5, 0, 0, 0], species=[Species(D=40, N=4, ExP=np.eye(4) * 0.3)])
    r = simulate(s)
    bg = (r.labels[0] == 255).sum()
    expected = 5e3 * 2.0
    assert abs(bg - expected) < 5 * np.sqrt(expected)
    assert r.header["MI_Bins"] == 4 * 2 ** 14 + 1   # as Sim.m in basic mode with noise


def test_advanced_mode_runs():
    sp = Species(D=40, N=4, color=2, fret=2, ExP=np.eye(4) * 0.3, aniso=np.tile([0.38, 1.0, 0.0, 1.0], (4, 1)))
    s = small(use_irf=True, use_aniso=True, use_fret_width=True, use_linker_width=True, species=[sp])
    r = simulate(s)
    assert r.mode == "advanced" and r.n_channels == 8
    par, perp = r.macrotimes[0].size, r.macrotimes[4].size
    assert 0 < perp < par
    for c in range(8):
        if r.microtimes[c].size:
            assert r.microtimes[c].max() < 8192


def test_dynamic_two_state():
    R1 = np.zeros((4, 4)); R1[1, 0] = 40
    R2 = np.zeros((4, 4)); R2[1, 0] = 90
    s = small(sim_time=1.0, dynamic_rates=np.array([[0, 1.0], [1.0, 0]]),
              species=[Species(D=40, N=4, color=2, fret=3, ExP=np.eye(4) * 0.3, R=R1),
                       Species(D=40, N=4, color=2, fret=3, ExP=np.eye(4) * 0.3, R=R2)])
    r = simulate(s)
    assert r.mode == "advanced" and r.macrotimes[0].size > 0 and r.macrotimes[1].size > 0


def test_raster_scan_image():
    s = SimSettings(box_size=[2000, 2000, 2000], sim_time=0.4, frames=4, scan_type=2, seed=1,
                    pixels=[20, 20], pixel_size=[50, 50], times=[250, 5, 0.1],
                    species=[Species(D=5, N=10, ExP=np.eye(4) * 0.3)])
    r = simulate(s)
    ticks_per_px = 250
    img = np.bincount((r.macrotimes[0] // ticks_per_px).astype(int) % 400, minlength=400)
    assert img.sum() == r.macrotimes[0].size > 0


# ---------------------------------------------------------------- correlation
def test_time_axis_like_matlab():
    t, e = time_axis(1e6)
    assert e == int(np.floor(np.log2(1e6 / 10)) + 1)
    np.testing.assert_array_equal(t[:22], np.r_[0:21, 22])
    assert t[-1] <= 1e6


def test_do_ccf_brute_force():
    rng = np.random.default_rng(0)
    t1 = np.sort(rng.integers(0, 5000, 300)); t2 = np.sort(rng.integers(0, 5000, 400))
    w1 = rng.random(300); w2 = rng.random(400)
    c = do_ccf(t1.astype(np.int64), t2.astype(np.int64), w1, w2, 10, 2)
    d = t2[None, :] - t1[:, None]
    ww = w1[:, None] * w2[None, :]
    for lag in range(0, 21):  # first two blocks are unbinned
        assert np.isclose(c[lag], ww[d == lag].sum())


def test_cross_correlation_uncorrelated_is_zero():
    rng = np.random.default_rng(1)
    T = 10_000_000.0
    d1 = [np.sort(rng.integers(0, int(T), 200_000)).astype(float)]
    d2 = [np.sort(rng.integers(0, int(T), 200_000)).astype(float)]
    g, lags = cross_correlation(d1, d2, T)
    sel = lags > 1000
    assert abs(g[sel, 0].mean()) < 0.01


@pytest.mark.skipif(not HAVE_COMPILERS, reason="needs gcc/g++ to build PAM's original code")
def test_correlator_identical_to_original_c_code():
    from validation.reference import matlab_cross_correlation, ref_do_ccf
    rng = np.random.default_rng(2)
    for _ in range(5):
        n1, n2 = rng.integers(10, 5000, 2)
        T = int(rng.integers(10_000, 5_000_000))
        t1 = np.sort(rng.integers(0, T, n1)); t2 = np.sort(rng.integers(0, T, n2))
        w1 = rng.normal(1, 0.3, n1); w2 = rng.normal(1, 0.3, n2)
        nb = int(np.floor(np.log2(T / 10)) + 1)
        np.testing.assert_array_equal(ref_do_ccf(t1, t2, w1, w2, 10, nb),
                                      do_ccf(t1.astype(np.int64), t2.astype(np.int64), w1, w2, 10, nb))
    M = 2_000_000.0
    D1 = [np.sort(rng.integers(0, int(M), 3000)).astype(float) for _ in range(3)]
    D2 = [np.sort(rng.integers(0, int(M), 4000)).astype(float) for _ in range(3)]
    ref, tref = matlab_cross_correlation(D1, D2, M)
    got, tgot = cross_correlation(D1, D2, M)
    np.testing.assert_array_equal(tref, tgot)
    np.testing.assert_allclose(got, ref, rtol=0, atol=1e-12)


@pytest.mark.skipif(not HAVE_COMPILERS, reason="needs g++ to build PAM's original kernels")
def test_matches_original_kernel_quick():
    from validation.reference import ref_simulate
    s = small(species=[Species(D=40, N=20, color=2, fret=2, ExP=np.eye(4) * 0.3, LT=[3, 2, 1, 1])])
    ref = [ref_simulate(s, seed=k) for k in range(6)]
    prt = []
    for k in range(6):
        s.seed = 50 + k
        r = simulate(s)
        prt.append((np.concatenate(r.macrotimes), np.concatenate(r.microtimes), np.concatenate(r.labels)))
    for code in (0b00000000, 0b00000101, 0b01010101):
        a = np.array([(x[2] == code).sum() for x in ref]) / s.sim_time
        b = np.array([(x[2] == code).sum() for x in prt]) / s.sim_time
        se = np.sqrt(a.var(ddof=1) / 6 + b.var(ddof=1) / 6)
        assert abs(a.mean() - b.mean()) < 4.5 * se
        ma = np.concatenate([x[1][x[2] == code] for x in ref]).astype(float)
        mb = np.concatenate([x[1][x[2] == code] for x in prt]).astype(float)
        assert abs(ma.mean() - mb.mean()) < 5 * np.sqrt(ma.var() / ma.size + mb.var() / mb.size)


# ---------------------------------------------------------------- I/O
def test_sim_file_roundtrip(tmp_path):
    s = small(species=[Species(D=40, N=4, ExP=np.eye(4) * 0.3)])
    r = simulate(s)
    files = r.save_sim(tmp_path / "test_1.sim")
    assert files == [str(tmp_path / "test_1.sim")]
    mts, mis, header = load_sim(files[0])
    assert len(mts) == 4
    np.testing.assert_array_equal(mts[0], r.macrotimes[0])
    np.testing.assert_array_equal(mis[0], r.microtimes[0])
    for key in ("Freq", "Frames", "FrameTime", "Lines", "Pixels", "MI_Bins"):  # read by LoadTcspc.m
        assert key in header
    assert header["Info"]["General"]["MIRange"] == 80.0
    from scipy.io import loadmat
    raw = loadmat(files[0])
    assert raw["Sim_Photons"].shape == (4, 2)
    assert raw["Sim_Photons"][0, 0].shape[1] == 1   # column vectors like MATLAB


def test_advanced_sim_writes_irf(tmp_path):
    sp = Species(D=40, N=2, ExP=np.eye(4) * 0.3)
    r = simulate(small(sim_time=0.05, use_irf=True, species=[sp]))
    files = r.save_sim(tmp_path / "adv_1.sim")
    assert files[1].endswith("_irf.sim") and os.path.exists(files[1])


def test_load_pam_json(tmp_path):
    pam = {"Name": "Sim A", "BS": [3000, 3000, 3000], "Seed": 5, "Freq": 2000, "SimTime": 0.1, "Frames": 1,
           "Px": [100, 100], "Size": [50, 50], "Time": [100, 10, 1], "ScanType": 1, "UseNoise": 0,
           "Noise": [0.2, 0.2, 0.2, 0.2], "MIRange": 50, "IRFWidth": 250, "HeterogeneityUpdate": 10,
           "LinkerWidth": 5, "DynamicRate": 0,
           "Species": [{"Name": "S1", "Color": 1, "FRET": 1, "Barrier": 1, "D": 30, "N": 3,
                        "wr": [200] * 4, "wz": [800] * 4, "dX": [0] * 4, "dY": [0] * 4, "dZ": [0] * 4,
                        "ExP": np.eye(4).tolist(), "Cross": np.eye(4).tolist(), "DetP": np.ones((4, 4)).tolist(),
                        "BlP": np.zeros((4, 4)).tolist(), "UseLT": 1, "UseIRF": 0, "LT": [3, 3, 3, 3],
                        "R0": np.zeros((4, 4)).tolist(), "R": np.zeros((4, 4)).tolist(),
                        "Aniso": np.tile([0.4, 1, 0, 1], (4, 1)).tolist(), "DistanceWidth": np.zeros((4, 4)).tolist()}]}
    p = tmp_path / "pam.json"
    p.write_text(json.dumps(pam))
    s = SimSettings.load_json(p)
    assert s.freq == 2000 and s.mi_range == 50 and s.species[0].D == 30 and s.species[0].wz[0] == 800
    r = simulate(s)
    assert r.header["Freq"] == 2e6
    s.save_json(tmp_path / "py.json")
    s2 = SimSettings.load_json(tmp_path / "py.json")
    assert s2.species[0].D == 30 and s2.freq == 2000


# ---------------------------------------------------------------- GUI helpers
def test_gui_helpers(tmp_path):
    pytest.importorskip("tkinter")
    from pamsim.gui import load_map_file, save_setup
    m = np.random.default_rng(0).random((50, 40))
    np.save(tmp_path / "m.npy", m)
    np.testing.assert_array_equal(load_map_file(str(tmp_path / "m.npy")), m)
    from scipy.io import savemat
    savemat(tmp_path / "m.mat", {"Map": np.dstack([m, m])})  # MATLAB: frames in 3rd dim
    assert load_map_file(str(tmp_path / "m.mat")).shape == (2, 50, 40)
    s = SimSettings(seed=None, maps=m, species=[Species(D=7, brightness=[50] * 4)])
    save_setup(s, tmp_path / "setup.json", seed=123)
    s2 = SimSettings.load_json(tmp_path / "setup.json")
    assert s2.seed == 123 and s2.species[0].D == 7 and s2.species[0].brightness[0] == 50 and s2.maps is None


def test_large_sim_is_split(tmp_path):
    for frames in (1, 6):
        s = small(frames=frames, species=[Species(D=40, N=10, color=2, ExP=np.eye(4) * 0.3)])
        r = simulate(s)
        n = sum(m.size for m in r.macrotimes)
        files = r.save_sim(tmp_path / f"big{frames}_1.sim", max_bytes=n * 10 / 2.5)  # forces 3 parts
        assert [os.path.basename(f) for f in files] == [f"big{frames}_1_part0{k}.sim" for k in (1, 2, 3)]
        offset, total = 0.0, 0.0
        got = [[] for _ in range(4)]
        for f in files:
            mts, mis, h = load_sim(f)
            span = h["Frames"] * h["FrameTime"]
            for c in range(4):
                assert np.all(mts[c] < span)
                got[c].append(mts[c] + offset)
            offset += span
        assert offset == r.header["Frames"] * r.header["FrameTime"]
        for c in range(4):
            np.testing.assert_array_equal(np.concatenate(got[c]), r.macrotimes[c])


def test_gui_presets_and_checks():
    pytest.importorskip("tkinter")
    from pamsim.gui import check_settings, presets
    for name, s in presets().items():
        s.sim_time = 0.05
        s.seed = 1  # fixed: the smFRET preset has only 3 molecules, which may all miss the focus in 50 ms
        msgs = check_settings(s)
        assert not [m for lvl, m in msgs if lvl == "error"], (name, msgs)
        r = simulate(s)
        assert sum(len(m) for m in r.macrotimes) > 0, name
    s = presets()["3-colour FCCS (PIE)"]
    s.species[0].brightness = np.array([10000.0, 100, 100, 0])
    assert any(lvl == "error" and "saturates" in m for lvl, m in check_settings(s))
    # the estimate predicts the simulated count rate, also when saturated
    s = SimSettings(box_size=[2000, 2000, 3000], sim_time=0.4, seed=1,
                    species=[Species(D=40, N=40, brightness=[300, 0, 0, 0])])
    est = float([m for lvl, m in check_settings(s) if "Expected" in m][0].split(":")[1].split("kHz")[0])
    got = simulate(s).count_rates_khz()[0]
    assert abs(got / est - 1) < 0.25


def test_read_planner():
    pytest.importorskip("openpyxl")
    from pamsim.planner import read_planner
    path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "planning", "PAMsim_simulation_planner.xlsx")
    rows = read_planner(path)
    ids = [r.sim_id for r in rows]
    assert ids[:3] == ["S00", "S01", "S02"] and "S11" in ids
    s00 = rows[0].settings
    assert rows[0].extras == "-" and not rows[0].direct_excitation  # crosstalk sheet: S00 listed, switched off
    assert len(s00.species) == 4 and s00.species[0].color == 3
    assert list(s00.species[2].brightness[:3]) == [0, 70, 50]
    fret = s00.species[3]
    assert fret.fret == 2 and fret.R[1, 0] == 50 and fret.R0[1, 0] == 63
    assert rows[0].settings.sim_time == 60 and rows[0].folder.endswith("Simulation")
    s11 = rows[ids.index("S11")]
    assert [sp.N for sp in s11.settings.species] == [200, 2000, 4000]
    s01 = rows[1].settings
    s01.sim_time = 0.05
    assert simulate(s01).macrotimes[0].size > 0


# ---------------------------------------------------------------- crosstalk / direct excitation sheet
PLANNER_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "..", "planning")


def test_planner_crosstalk_sheet(tmp_path):
    pytest.importorskip("openpyxl")
    import subprocess
    from openpyxl import load_workbook
    from pamsim.planner import read_planner
    path = tmp_path / "planner.xlsx"
    subprocess.run([sys.executable, os.path.join(PLANNER_DIR, "make_planner.py"), "-o", str(path)], check=True,
                   capture_output=True)
    rows = {r.sim_id: r for r in read_planner(path)}
    assert rows["S00"].extras == "-" and not rows["S00"].crosstalk  # example values are there but switched off
    assert np.all(rows["S00"].settings.species[0].Cross == np.eye(4))
    wb = load_workbook(path)
    xs = wb["Crosstalk & direct exc"]
    assert xs["A4"].value == "S00"
    xs["D4"] = "Yes"; xs["H4"] = "Yes"                       # S00: both on (0.08/0.01/0.05, 0.05/0.01/0.06)
    xs["A5"] = "S01"; xs["H5"] = "Yes"; xs["I5"] = 0.2      # S01: direct excitation only
    xs["A6"] = "S02"; xs["D6"] = "No"; xs["E6"] = 0.3       # S02: listed but switched off
    wb.save(path)
    rows = {r.sim_id: r for r in read_planner(path)}
    s00 = rows["S00"]
    assert s00.extras == "crosstalk + direct exc." and s00.crosstalk == {"1→2": 0.08, "1→3": 0.01, "2→3": 0.05}
    freq = s00.settings.freq
    for sp in s00.settings.species:
        np.testing.assert_allclose(sp.Cross[1:3, 0], [0.08, 0.01]); assert sp.Cross[2, 1] == 0.05
        e = sp.excitation_matrix(freq)
        np.testing.assert_allclose(np.diag(e)[:3], sp.brightness[:3] / freq * sp.Cross[:, :3].sum(axis=0))
        np.testing.assert_allclose([e[1, 0], e[2, 0], e[2, 1]], [0.05 * e[1, 1], 0.01 * e[2, 2], 0.06 * e[2, 2]])
        np.testing.assert_allclose(sp.excitation_matrix(freq), e)  # stable when applied again
    assert rows["S01"].extras == "direct exc." and rows["S01"].settings.species[0].Cross[1, 0] == 0
    assert rows["S02"].extras == "-" and rows["S03"].extras == "-"


def test_crosstalk_and_direct_excitation_simulated():
    from pamsim.planner import apply_crosstalk
    s = small(sim_time=1.0, frames=1, species=[Species(name="dye 1", D=20, N=60, color=2, brightness=[100, 0, 0, 0]),
                                                Species(name="dye 2", D=20, N=60, color=2, brightness=[0, 100, 0, 0])])
    for sp in s.species:
        apply_crosstalk(sp, s.freq, {"1→2": 0.1, "2→3": 0.5}, {"1→2": 0.2})  # 2→3 ignored: only 2 colours
        assert sp.Cross[2, 1] == 0
    r = simulate(s)
    lab = np.concatenate(r.labels)
    laser, exc, emit, chan = lab >> 6, (lab >> 4) & 3, (lab >> 2) & 3, lab & 3
    d1 = emit == 0
    assert abs(np.sum(d1 & (chan == 1)) / np.sum(d1 & (chan == 0)) / 0.1 - 1) < 0.15   # crosstalk dye 1 -> ch 2
    d2 = (emit == 1) & (chan == 1)
    assert abs(np.sum(d2 & (laser == 0)) / np.sum(d2 & (laser == 1)) / 0.2 - 1) < 0.15  # laser 1 excites dye 2
    assert np.all(exc == emit) and not np.any(chan >= 2)


def test_info_report(tmp_path):
    from pamsim.report import simulation_info
    R = np.zeros((4, 4)); R[1, 0] = 50
    R0 = np.zeros((4, 4)); R0[1, 0] = 63
    s = small(name="demo", sim_time=60, mi_range=50, use_noise=True, noise=[0.2] * 4, species=[
        Species(name="green only", D=30, N=100, color=2, brightness=[80, 0, 0, 0], LT=[4, 3, 2, 2]),
        Species(name="FRET pair", D=50, N=20, color=2, fret=2, brightness=[100, 70, 0, 0], LT=[4, 3, 2, 2], R=R, R0=R0)])
    text = simulation_info(s, file_name="demo_1.sim", seed=7, sim_id="S12")
    lines = [ln.strip() for ln in text.splitlines()]
    assert lines[2:7] == ["File name:                   demo_1.sim".strip(), "Duration:                    60 s",
                          "Microtime range:             50 ns", "Background:                  yes (0.2 kHz per channel)",
                          "Number of species:           2"]
    assert "FRET:                        static" in lines
    assert any(ln.startswith("R0:") and "1->2 = 63" in ln for ln in lines)
    assert any(ln.startswith("R (FRET pair):") and "1->2 = 50 (E = 0.800)" in ln for ln in lines)
    assert "Crosstalk:                   no" in lines and "Direct excitation:           no" in lines
    sp1 = lines[lines.index("Species 1"):lines.index("Species 2")]
    assert "Number of colours:         1 (colour 1)" in sp1 and "Brightness:                80 kHz" in sp1
    assert "Diffusion coefficient:     30 um^2/s" in sp1 and "Lifetime:                  4 ns" in sp1
    sp2 = lines[lines.index("Species 2"):lines.index("Run")]
    for k, b, lt in ((1, 100, 4), (2, 70, 3)):
        assert f"Diffusion coefficient {k}:   50 um^2/s" in sp2 and f"Brightness {k}:              {b} kHz" in sp2
        assert f"Lifetime {k}:                {lt} ns" in sp2 and any(ln.startswith(f"Concentration {k}:") for ln in sp2)
    assert "Planner Sim ID:            S12" in lines and "Seed:                      7" in lines
    # with crosstalk / direct excitation, and written next to a simulated file
    from pamsim.planner import apply_crosstalk
    for sp in s.species:
        apply_crosstalk(sp, s.freq, {"1→2": 0.05}, {"1→2": 0.1})
    s.sim_time = 0.2
    r = simulate(s)
    files = r.save_sim(tmp_path / "demo_1.sim")
    path = r.save_info(tmp_path / "demo_1_info.txt", files=files)
    text = open(path, encoding="utf-8").read()
    assert "File name:                   demo_1.sim" in text and "Duration:                    0.2 s" in text
    assert "Crosstalk:                   yes (dye 1 -> channel 2 = 0.05)" in text
    assert "Direct excitation:           yes (laser 1 -> dye 2 = 0.1)" in text


def test_check_rates_include_crosstalk():
    pytest.importorskip("tkinter")
    from pamsim.gui import check_settings
    from pamsim.planner import apply_crosstalk

    def expected(s):
        msg = [m for lvl, m in check_settings(s) if "Expected" in m][0]
        return np.array([float(x) for x in msg.split(":")[1].split("kHz")[0].split(",")])
    s = small(sim_time=1.0, frames=1, species=[Species(D=20, N=60, color=2, brightness=[100, 0, 0, 0]),
                                                Species(D=20, N=60, color=2, brightness=[0, 100, 0, 0])])
    before = expected(s)
    for sp in s.species:
        apply_crosstalk(sp, s.freq, {"1→2": 0.1}, {"1→2": 0.2})
    est = expected(s)
    assert abs(est[0] / before[0] - 1) < 0.01 and est[1] > before[1] * 1.25
    got = np.array(simulate(s).count_rates_khz()[:2])
    assert np.all(np.abs(got / est - 1) < 0.2), (got, est)


def test_planner_migration(tmp_path):
    pytest.importorskip("openpyxl")
    import subprocess
    from openpyxl import load_workbook
    from pamsim.planner import read_planner
    make = [sys.executable, os.path.join(PLANNER_DIR, "make_planner.py")]
    old, new = tmp_path / "old.xlsx", tmp_path / "new.xlsx"
    subprocess.run(make + ["-o", str(old)], check=True, capture_output=True)
    wb = load_workbook(old)
    wb["Crosstalk & direct exc"]["D4"] = "Yes"
    wb["Simulations"]["B5"] = "renamed"
    wb.save(old)
    subprocess.run(make + ["--from", str(old), "--species", "2", "-o", str(new)], check=True, capture_output=True)
    a, b = read_planner(old), read_planner(new)
    assert [r.sim_id for r in a] == [r.sim_id for r in b] and b[1].settings.name == "renamed"
    for ra, rb in zip(a, b):
        assert ra.settings.to_dict() == rb.settings.to_dict() and ra.crosstalk == rb.crosstalk
        assert (ra.folder, ra.purpose, ra.nfiles) == (rb.folder, rb.purpose, rb.nfiles)
    assert b[0].extras == "crosstalk"
