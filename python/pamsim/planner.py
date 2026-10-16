"""Read simulations from the Excel planner (planning/PAMsim_simulation_planner.xlsx).

The planner has one row per simulation on the sheet "Simulations": band
titles in row 1 ("Measurement", "Species 1", ...), column headers in row 3,
data from row 4.  Columns are found by their header text, so the planner
may be rearranged or extended (more species blocks) without breaking this.

The optional sheet "Crosstalk & direct exc" has the same layout and one row
per simulation (matched by Sim ID).  Its values are only used when the row's
switch "Crosstalk?" / "Direct exc.?" is Yes; planners without this sheet
simulate without crosstalk and direct excitation, as before.
"""

from __future__ import annotations

import numpy as np

from .params import SimSettings, Species

SCAN = {"point": 1, "raster scan": 2, "line scan": 3, "circle scan": 4}
FRET = {"no fret": 1, "static fret": 2, "dynamic fret": 3}
FORMAT = {".sim": 0, ".npz": 1, "both": 2}
#: (bluer colour, redder colour) pairs of the crosstalk sheet, 0-based
PAIRS = {"1→2": (0, 1), "1→3": (0, 2), "2→3": (1, 2)}


class PlannerRow:
    """One simulation of the planner."""

    def __init__(self, sim_id, settings, folder, nfiles, fmt, purpose, status, check,
                 crosstalk=None, direct_excitation=None):
        self.sim_id = sim_id
        self.settings = settings
        self.folder = folder
        self.nfiles = nfiles
        self.fmt = fmt
        self.purpose = purpose
        self.status = status
        self.check = check
        #: values in effect, e.g. {"1→2": 0.05}; empty = switched off
        self.crosstalk = crosstalk or {}
        self.direct_excitation = direct_excitation or {}

    @property
    def extras(self):
        """Short label: which of crosstalk / direct excitation are switched on."""
        on = [n for n, d in (("crosstalk", self.crosstalk), ("direct exc.", self.direct_excitation)) if d]
        return " + ".join(on) if on else "-"

    def __repr__(self):
        return f"PlannerRow({self.sim_id!r}, {self.settings.name!r}, {len(self.settings.species)} species)"


def _num(v, default=None):
    if v is None or (isinstance(v, str) and not v.strip()):
        return default
    try:
        return float(v)
    except (TypeError, ValueError):
        return default


def _text(v):
    return "" if v is None else str(v).strip()


def _columns(ws):
    """Map (band, header) -> column index; band = merged title of row 1."""
    band = ""
    cols = {}
    for c in range(1, ws.max_column + 1):
        b = ws.cell(row=1, column=c).value
        if b not in (None, ""):
            band = str(b).strip()
        h = ws.cell(row=3, column=c).value
        if h not in (None, ""):
            cols[(band, str(h).strip())] = c
    return cols


def _read_crosstalk_sheet(wb):
    """Sim ID -> (crosstalk, direct excitation) from the optional crosstalk sheet.

    Each is a dict like {"1→2": 0.05} with the pairs that are switched on and
    non-zero; a switch that is not "Yes" gives an empty dict.  If a Sim ID is
    listed twice, the first row counts (as in the planner's formulas).
    """
    name = next((n for n in wb.sheetnames if n.lower().startswith("crosstalk")), None)
    if name is None:
        return {}
    ws = wb[name]
    cols = {head: c for (_, head), c in _columns(ws).items()}
    if "Sim ID" not in cols:
        return {}
    out = {}
    for vals in ws.iter_rows(min_row=4, values_only=True):
        def g(head):
            c = cols.get(head)
            return None if c is None or c - 1 >= len(vals) else vals[c - 1]

        sim_id = _text(g("Sim ID"))
        if not sim_id or sim_id in out:
            continue
        found = []
        for switch, prefix in (("Crosstalk?", "Crosstalk"), ("Direct exc.?", "Direct exc.")):
            on = _text(g(switch)).lower() == "yes"
            values = {pair: _num(g(f"{prefix} {pair}"), 0.0) or 0.0 for pair in PAIRS} if on else {}
            found.append({pair: v for pair, v in values.items() if v > 0})
        out[sim_id] = tuple(found)
    return out


def apply_crosstalk(species, freq, crosstalk=None, direct_excitation=None):
    """Put crosstalk and direct excitation (fractions, keys "1→2", "1→3", "2→3") into a species.

    * crosstalk "1→2" = photons of dye 1 in channel 2 relative to channel 1
      (``Cross[1, 0]``, Sim.m's relative crosstalk);
    * direct excitation "1→2" = excitation of dye 2 by laser 1 relative to its
      excitation by its own laser 2 (``ExP[1, 0] = value * ExP[1, 1]``).

    The brightness keeps its meaning (count rate of a dye in its own channel).
    ``ExP`` is stored with the diagonal that this brightness gives, so
    :meth:`Species.excitation_matrix` leaves the direct excitation unchanged,
    also for species without the exciting laser's own dye.  Pairs with a colour
    above ``species.color`` are ignored.
    """
    cross = np.eye(4)
    for pair, v in (crosstalk or {}).items():
        dye, chan = PAIRS[pair]
        if chan < species.color:
            cross[chan, dye] = v
    species.Cross = cross
    species.ExP = np.eye(4)
    exp_ = species.excitation_matrix(freq)  # diagonal from the brightness (incl. the crosstalk)
    for pair, v in (direct_excitation or {}).items():
        laser, dye = PAIRS[pair]
        if dye < species.color:
            exp_[dye, laser] = v * exp_[dye, dye]
    species.ExP = exp_
    return species


def read_planner(path):
    """Return a list of :class:`PlannerRow`, one per filled simulation row."""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    if "Simulations" not in wb.sheetnames:
        raise ValueError("this file has no 'Simulations' sheet - is it a PAMsim planner?")
    extra = _read_crosstalk_sheet(wb)
    ws = wb["Simulations"]
    cols = _columns(ws)
    by_header = {}
    for (band, head), c in cols.items():
        if not band.startswith("Species"):
            by_header.setdefault(head, c)
    species_bands = sorted({band for band, _ in cols if band.startswith("Species")},
                           key=lambda b: int(b.split()[-1]) if b.split()[-1].isdigit() else 0)
    if "Sim ID" not in by_header:
        raise ValueError("no 'Sim ID' column found in row 3 of the Simulations sheet")
    rows = list(ws.iter_rows(min_row=4, values_only=True))
    wb.close()

    out = []
    for vals in rows:
        def g(head, default=None, band=None):
            c = cols.get((band, head)) if band else by_header.get(head)
            if c is None or c - 1 >= len(vals):
                return default
            v = vals[c - 1]
            return default if v is None else v

        sim_id = _text(g("Sim ID"))
        if not sim_id:
            continue
        ncol = int(_num(g("Colours"), 1) or 1)
        ncol = max(1, min(4, ncol))
        # focus per colour (laser); older planners had one wr / wz for all colours
        wr_all = _num(g("wr [nm]"), 200.0)
        wz_all = _num(g("wz [nm]"), 1000.0)
        wr = [_num(g(f"wr {k} [nm]"), wr_all) for k in (1, 2, 3)]
        wz = [_num(g(f"wz {k} [nm]"), wz_all) for k in (1, 2, 3)]
        shift = {ax: [_num(g(f"{ax} {k} [nm]"), 0.0) for k in (1, 2, 3)] for ax in ("dX", "dY", "dZ")}
        r0 = np.zeros((4, 4))
        r0[1, 0] = _num(g("R0 1→2"), 63.0)
        r0[2, 0] = _num(g("R0 1→3"), 51.0)
        r0[2, 1] = _num(g("R0 2→3"), 68.0)
        species = []
        for band in species_bands:
            name = _text(g("Name", band=band))
            n = _num(g("Molecules N", band=band), 0) or 0
            if not name and n <= 0:
                continue
            bright = [_num(g(f"Bright. {k} [kHz]", band=band), 0.0) or 0.0 for k in (1, 2, 3)]
            lt = [_num(g(f"Lifetime {k} [ns]", band=band), 4.0) for k in (1, 2, 3)]
            fret = FRET.get(_text(g("FRET", band=band)).lower(), 1)
            # FRET only between dyes that exist (brightness > 0) and pairs with a distance R;
            # an empty R (or a missing dye) switches that pair off (R0 = 0 -> rate 0)
            R = np.zeros((4, 4))
            sp_r0 = r0.copy()
            for (a, d), head in (((1, 0), "R 1→2"), ((2, 0), "R 1→3"), ((2, 1), "R 2→3")):
                dist = _num(g(head, band=band))
                if dist is None or dist <= 0 or bright[a] <= 0 or bright[d] <= 0:
                    sp_r0[a, d] = 0.0
                    R[a, d] = 1e4
                else:
                    R[a, d] = dist
            species.append(Species(
                name=name or band, color=ncol, fret=fret, D=_num(g("D [µm²/s]", band=band), 1.0) or 0.0,
                N=int(round(n)), wr=wr + [wr[-1]], wz=wz + [wz[-1]],
                dX=shift["dX"] + [0.0], dY=shift["dY"] + [0.0], dZ=shift["dZ"] + [0.0], brightness=bright + [0.0],
                LT=lt + [lt[-1]], R0=sp_r0, R=R))
        seed = _num(g("Seed"))
        bgk = _num(g("Background [kHz/ch]"), 0.0) or 0.0
        settings = SimSettings(
            name=_text(g("File name")) or sim_id,
            box_size=[_num(g("Box x [nm]"), 5000.0), _num(g("Box y [nm]"), 5000.0), _num(g("Box z [nm]"), 5000.0)],
            seed=int(seed) if seed is not None else None,
            freq=_num(g("Sampling freq. [kHz]"), 1000.0),
            sim_time=_num(g("Duration [s]"), 10.0),
            frames=int(_num(g("Frames"), 1) or 1),
            scan_type=SCAN.get(_text(g("Scan type")).lower(), 1),
            mi_range=_num(g("Microtime range [ns]"), 80.0),
            use_noise=_text(g("Background?")).lower() == "yes",
            noise=[bgk] * 4,
            use_irf=_text(g("IRF?")).lower() == "yes",
            irf_width=_num(g("IRF width [ps]"), 250.0),
            species=species or [Species(color=ncol, wr=wr + [wr[-1]], wz=wz + [wz[-1]])],
        )
        # crosstalk / direct excitation (optional sheet); only pairs within the colours of the setup
        ct, de = extra.get(sim_id, ({}, {}))
        ct = {k: v for k, v in ct.items() if PAIRS[k][1] < ncol}
        de = {k: v for k, v in de.items() if PAIRS[k][1] < ncol}
        if ct or de:
            for sp in settings.species:
                apply_crosstalk(sp, settings.freq, ct, de)
        out.append(PlannerRow(
            sim_id=sim_id, settings=settings, folder=_text(g("Output folder")),
            nfiles=int(_num(g("Number of files"), 1) or 1), fmt=FORMAT.get(_text(g("Format")).lower(), 0),
            purpose=_text(g("Purpose / question")), status=_text(g("Status")),
            check=_text(g("Check")), crosstalk=ct, direct_excitation=de,
        ))
    return out
