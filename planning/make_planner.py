"""Generate a PAMsim simulation planner (Excel), one row per simulation.

Usage (in a terminal, inside the "planning" folder):

    python make_planner.py                         # new planner with the example rows
    python make_planner.py --empty                 # new, empty planner
    python make_planner.py --species 10            # 10 species blocks per simulation instead of 6
    python make_planner.py -o my_project.xlsx      # choose the file name
    python make_planner.py --rows 200              # prepare 200 simulation rows instead of 60
    python make_planner.py --from old.xlsx -o new.xlsx   # new layout, your simulations copied from old.xlsx

An existing file is never overwritten unless you add --force.
Needs:  python -m pip install openpyxl
"""
import os

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation

import argparse
import sys

_ap = argparse.ArgumentParser(description="Generate a PAMsim simulation planner (Excel).")
_ap.add_argument("-o", "--output", default=os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                                         "PAMsim_simulation_planner.xlsx"),
                 help="output file (default: PAMsim_simulation_planner.xlsx next to this script)")
_ap.add_argument("--species", type=int, default=6, help="species blocks per simulation (default 6)")
_ap.add_argument("--rows", type=int, default=60, help="simulation rows prepared (default 60)")
_ap.add_argument("--empty", action="store_true", help="no example rows")
_ap.add_argument("--force", action="store_true", help="overwrite an existing file")
_ap.add_argument("--from", dest="source", metavar="OLD.xlsx",
                 help="copy all input cells (simulations, species, crosstalk sheet) from an existing planner "
                      "into the new layout instead of the example rows")
ARGS = _ap.parse_args()
if os.path.exists(ARGS.output) and not ARGS.force:
    sys.exit(f"{ARGS.output} already exists - choose another name with -o, or add --force to overwrite it.")
if ARGS.source and os.path.abspath(ARGS.source) == os.path.abspath(ARGS.output):
    sys.exit("--from and -o must be different files (the old planner is never changed).")


def read_old_planner(path):
    """Input values of an existing planner: {"sim": [row, ...], "xs": [row, ...], "species": n}.

    A row maps (band title, header) -> value; formulas are skipped.  Rows without a Sim ID are left out.
    """
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True)
    out = {"sim": [], "xs": [], "species": 0}
    for key, name in (("sim", "Simulations"), ("xs", next((n for n in wb.sheetnames
                                                           if n.lower().startswith("crosstalk")), None))):
        if name not in wb.sheetnames:
            continue
        rows = list(wb[name].iter_rows(values_only=True))
        if len(rows) < 4:
            continue
        band, heads = "", []
        for b, hd in zip(rows[0], rows[2]):
            band = str(b).strip() if b not in (None, "") else band
            heads.append((band, str(hd).strip()) if hd not in (None, "") else None)
        if key == "sim":
            out["species"] = len({b for b, _ in filter(None, heads) if b.startswith("Species ")})
        for vals in rows[3:]:
            rec = {hd: v for hd, v in zip(heads, vals)
                   if hd and v not in (None, "") and not (isinstance(v, str) and v.startswith("="))}
            if any(h == "Sim ID" for _, h in rec):
                out[key].append(rec)
    wb.close()
    return out


OLD = read_old_planner(ARGS.source) if ARGS.source else None
MAX_SPECIES = max(1, ARGS.species, OLD["species"] if OLD else 0)
N_ROWS = max(1, ARGS.rows, len(OLD["sim"]) + 20 if OLD else 0)  # simulation rows prepared
FIRST, LAST = 4, 4 + N_ROWS - 1  # data rows (rows 1-3: bands and headers)
OUT = ARGS.output
XS_NAME = "Crosstalk & direct exc"  # optional sheet: crosstalk and direct excitation per simulation
XS = f"'{XS_NAME}'!"

F = "Arial"
HDR = PatternFill("solid", fgColor="1F3A5F")
INPUT = PatternFill("solid", fgColor="FFF6C7")
CALC = PatternFill("solid", fgColor="EEF2F6")
thin = Side(style="thin", color="C9CED6")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
BLUE = Font(name=F, size=10, color="0000FF")
BLACK = Font(name=F, size=10, color="000000")
WHITE_B = Font(name=F, size=10, bold=True, color="FFFFFF")
BAND = {"id": "2A78D6", "gen": "3D6B9E", "foc": "5A7FA8", "fret": "7D5BA6", "out": "6C8EAF",
        "chk": "6E6E6E", "res": "4E8A5B"}
SP_BAND = ["1B6F8A", "8A5A1B", "2E7D4F", "7A2E5E", "4F4F8A", "6B6B1B", "1B6B6B", "8A1B1B"]
RED = PatternFill("solid", fgColor="F8D3D3")
ORANGE = PatternFill("solid", fgColor="FCE6C4")
GREEN = PatternFill("solid", fgColor="D5EFD5")

wb = Workbook()

# ====================================================================== How to use
h = wb.active
h.title = "How to use"
lines = [
    ("PAMsim simulation planner", Font(name=F, size=16, bold=True, color="1F3A5F")),
    ("Plan your simulations here, then type the values into PAMsim (or Sim.m).", Font(name=F, size=10, italic=True, color="555555")),
    ("", None),
    ("How it works", Font(name=F, size=12, bold=True)),
    ("•  Sheet 'Simulations': ONE ROW PER SIMULATION. Left: identification and measurement settings.", None),
    (f"•  Then {MAX_SPECIES} species blocks side by side (Species 1, Species 2, ...). Each species has its own name, D, number of", None),
    ("    molecules N, brightness and lifetime per colour, and FRET (type and dye distances R). Leave unused species blocks empty.", None),
    ("•  Brightness 0 = this species carries no dye of that colour (e.g. a green-red complex: B1 = 0, B2 = 70, B3 = 50).", None),
    ("•  Focus per colour: each colour (laser) has its own focus size wr / wz and position dX / dY / dZ (0 = foci overlap).", None),
    ("•  FRET: Förster radii R0 are set once per simulation (purple block); per species choose FRET type and distances R.", None),
    ("    The FRET efficiency E = 1 / (1 + (R/R0)^6) is calculated for you (grey cells).", None),
    ("•  Grey cells are calculated: concentration, FRET efficiency, a check per species, expected count rates, photons, memory.", None),
    ("•  Use the + / - buttons above the columns (Excel grouping) to fold species blocks away.", None),
    ("•  Fix every red check before running a simulation; orange checks bias the results slightly.", None),
    (f"•  Optional: crosstalk and direct excitation are on the separate sheet '{XS_NAME}' (off unless you switch them on).", None),
    ("    One row per simulation: pick its Sim ID, set 'Crosstalk?' and/or 'Direct exc.?' to Yes and give the values as fractions.", None),
    ("    Simulations that are not listed there (or have both switches on No) are simulated without crosstalk and direct excitation.", None),
    ("•  Empty cells are treated like in PAMsim: box 5000 nm, 1000 kHz, 10 s, 1 colour, wr 200 nm, wz 1000 nm, R0 63/51/68.", None),
    ("", None),
    ("Colour code", Font(name=F, size=12, bold=True)),
    ("Yellow cells with blue text = your input.   Grey cells with black text = formulas (do not overwrite).", None),
    ("", None),
    ("Rules of thumb (from the PAMsim checks)", Font(name=F, size=12, bold=True)),
    ("•  Brightness < 10% of the sampling frequency (e.g. < 100 kHz at 1000 kHz). Above it the focus saturates:", None),
    ("    the fitted D comes out too small and the concentration too large. At >= 100% the excitation is fully saturated.", None),
    ("•  One diffusion step sqrt(2·D·dt) should be < 10% of wr; otherwise raise the sampling frequency.", None),
    ("•  Box >= 8·wr laterally and >= 3·wz axially.", None),
    ("•  Memory ≈ 12 bytes per photon and file. For very long runs use several files (Number of files) instead.", None),
    ("•  Expected count rates: Gaussian focus incl. saturation (as in PAMsim's Check panel), plus crosstalk and direct excitation", None),
    ("    when they are switched on. FRET is ignored, so with FRET the donor channel gets fewer and the acceptor channel more photons.", None),
    ("", None),
    ("Units", Font(name=F, size=12, bold=True)),
    ("nm (box, wr, wz) · kHz (sampling frequency, brightness, background) · s (duration) · ns (lifetime, microtime range)", None),
    ("µm²/s (D) · nM (concentration) · Å or nm for R0 and R (use the same unit for both) · kHz (count rates) · million photons · GB", None),
    ("Crosstalk and direct excitation: fractions without unit (0.05 = 5 %).", None),
    ("", None),
    ("Example rows", Font(name=F, size=12, bold=True)),
    ("S00 is a worked example (realistic 3-colour FCCS mixture with a FRET pair). S01-S11 are the simulations from your earlier", None),
    ("spreadsheet (box 5000 nm, wr 200 nm, wz 1000 nm for all colours assumed) - the checks show which of them saturate.", None),
    ("S00 also shows colour-dependent foci (wr 180/210/240 nm, wz 900/1050/1200 nm, colour 3 shifted 100 nm in z).", None),
    (f"More than {MAX_SPECIES} species? Run  python make_planner.py --species 10 -o new_planner.xlsx  (see the README).", None),
]
for i, (t, f) in enumerate(lines, start=1):
    c = h.cell(row=i, column=2, value=t)
    c.font = f or Font(name=F, size=10)
h.column_dimensions["A"].width = 3
h.column_dimensions["B"].width = 125
h.cell(row=next(i for i, (t, _) in enumerate(lines, start=1) if t.startswith("Yellow cells")), column=2).fill = INPUT
h.sheet_view.showGridLines = False

# ====================================================================== Column guide
GUIDE = [
    ("Identification", None, None, None, None),
    ("Sim ID", "Unique short name of the simulation (e.g. S12)", "-", "Required", "row is skipped"),
    ("File name", "Name of the simulation: folder <Output folder>/<name>/ with files <name>_1.sim, ...", "-", "Optional", "the Sim ID"),
    ("Purpose / question", "What you want to learn from this simulation", "-", "Not used", "-"),
    ("Output folder", "Main folder; each simulation gets its own subfolder (created automatically)", "-", "Optional", "PAMsim's output folder"),
    ("Status, Priority", "Your bookkeeping", "-", "Not used", "-"),
    ("Measurement", None, None, None, None),
    ("Box x, y, z", "Size of the simulation box; z = 0 gives 2D (membrane)", "nm", "Optional", "5000 each"),
    ("Sampling freq.", "Time resolution of the random walk (time step = 1/frequency)", "kHz", "Optional", "1000"),
    ("Duration", "Simulated measurement time per file", "s", "Required", "10"),
    ("Frames", "Number of frames (relevant for scanning)", "-", "Optional", "1"),
    ("Microtime range", "TCSPC range (PIE period)", "ns", "Optional", "80"),
    ("Scan type", "Point / Raster scan / Line scan / Circle scan", "-", "Optional", "Point"),
    ("Colours", "Number of colours = PIE lasers = detection channels (1-3)", "-", "Required", "1 (brightness 2/3 ignored!)"),
    ("Background? / Background", "Uncorrelated background and its count rate per channel", "kHz", "Optional", "No / 0"),
    ("IRF? / IRF width", "Simulate an instrument response function", "ps", "Optional", "No / 250"),
    ("Focus per colour and Förster radii", None, None, None, None),
    ("wr 1/2/3", "Lateral 1/e² radius of the focus of each colour (laser)", "nm", "Optional", "200"),
    ("wz 1/2/3", "Axial 1/e² radius of the focus of each colour (laser)", "nm", "Optional", "1000"),
    ("dX, dY, dZ 1/2/3", "Shift of each colour's focus from the box centre (focus overlap for FCCS)", "nm", "Optional", "0"),
    ("R0 1→2, 1→3, 2→3", "Förster radius per dye pair (donor colour → acceptor colour)", "Å or nm", "Optional", "63, 51, 68"),
    ("Species blocks (one block per species; leave unused blocks empty)", None, None, None, None),
    ("Name", "Name of the species", "-", "Optional", "Species 1, 2, ..."),
    ("D", "Diffusion coefficient", "µm²/s", "Required", "1"),
    ("Molecules N", "Number of molecules in the box", "-", "Required", "0 = species absent"),
    ("Brightness 1/2/3", "Count rate of one molecule in the focus centre per colour; 0 = no dye of this colour", "kHz", "Required", "0"),
    ("Lifetime 1/2/3", "Fluorescence lifetime per colour", "ns", "Optional", "4"),
    ("FRET", "No FRET / Static FRET / Dynamic FRET", "-", "Optional", "No FRET"),
    ("R 1→2, 1→3, 2→3", "Dye distance for FRET between two colours; only dyes with brightness > 0 take part", "Å or nm", "Optional", "no FRET for this pair"),
    ("E, Conc., Check (grey)", "FRET efficiency 1/(1+(R/R0)^6), concentration, check of the species", "-, nM", "Calculated", "-"),
    ("Totals, output and results", None, None, None, None),
    ("Exp. rates, Photons, Memory (grey)", "Expected count rate per channel (incl. saturation, crosstalk and direct excitation), photons and memory per file", "kHz, M, GB", "Calculated", "-"),
    ("Crosstalk / dir. exc. (grey)", f"Shows whether crosstalk and/or direct excitation are switched on for this simulation (sheet '{XS_NAME}')", "-", "Calculated", "-"),
    ("Check (grey)", "Overall check of the simulation", "-", "Calculated", "-"),
    ("Number of files", "Independent repeats with the same parameters", "-", "Optional", "1"),
    ("Seed", "Random seed; file k uses seed + k - 1", "-", "Optional", "random"),
    ("Format", ".sim (PAM) / .npz / both", "-", "Optional", ".sim"),
    ("Date run ... Result notes", "Your notes after running and analysing", "-", "Not used", "-"),
    (f"Sheet '{XS_NAME}' (optional - one row per simulation; not listed = off)", None, None, None, None),
    ("Sim ID", "Sim ID of a simulation on the Simulations sheet (drop-down). File name and Colours are looked up (grey)", "-", "Required", "row is ignored"),
    ("Crosstalk?", "Yes = use the crosstalk values of this row", "-", "Optional", "No (no crosstalk)"),
    ("Crosstalk 1→2, 1→3, 2→3", "Photons of dye 1 (2) detected in channel 2 (3), as a fraction of that dye's photons in its own channel", "fraction", "Optional", "0"),
    ("Direct exc.?", "Yes = use the direct excitation values of this row", "-", "Optional", "No (no direct excitation)"),
    ("Direct exc. 1→2, 1→3, 2→3", "Excitation of dye 2 (3) by laser 1 (2), as a fraction of its excitation by its own laser", "fraction", "Optional", "0"),
    ("Check (grey)", "on / off, and warnings (unknown Sim ID, too few colours, values > 1, ...)", "-", "Calculated", "-"),
]
gs = wb.create_sheet("Column guide")
for j, (hd, w) in enumerate((("Column", 34), ("Meaning", 70), ("Unit", 11), ("Required?", 12), ("If left blank", 26)), start=1):
    c = gs.cell(row=1, column=j, value=hd)
    c.fill = HDR; c.font = WHITE_B; c.border = BOX
    gs.column_dimensions[L(j)].width = w
for i, row in enumerate(GUIDE, start=2):
    if row[1] is None:
        gs.merge_cells(start_row=i, start_column=1, end_row=i, end_column=5)
        c = gs.cell(row=i, column=1, value=row[0])
        c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="5A7FA8")
        continue
    for j, v in enumerate(row, start=1):
        c = gs.cell(row=i, column=j, value=v)
        c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(name=F, size=10, bold=(j == 4 and v == "Required"), color="9C1C1C" if (j == 4 and v == "Required") else "000000")
        if j == 4 and v == "Required":
            c.fill = PatternFill("solid", fgColor="FDE2E2")
gs.freeze_panes = "A2"
gs.cell(row=len(GUIDE) + 3, column=1, value="Not in the planner (set them in PAMsim after 'Load into editor'): "
        "detection efficiency, bleaching, anisotropy, dynamic FRET rates, barrier maps.").font = Font(name=F, italic=True, size=10)
gs.cell(row=len(GUIDE) + 4, column=1, value="Full explanation: planning/PAMsim_planner_guide.docx").font = Font(name=F, italic=True, size=10)

# ====================================================================== Lists
ls = wb.create_sheet("Lists")
for col, (title, vals) in {"A": ("Status", ["Idea", "Planned", "Running", "Done", "Analysed", "Discarded"]),
                           "B": ("Scan type", ["Point", "Raster scan", "Line scan", "Circle scan"]),
                           "C": ("Yes/No", ["Yes", "No"]),
                           "D": ("FRET", ["No FRET", "Static FRET", "Dynamic FRET"]),
                           "E": ("Format", [".sim", ".npz", "both"])}.items():
    ls[f"{col}1"] = title
    for i, v in enumerate(vals, start=2):
        ls[f"{col}{i}"] = v
ls.sheet_state = "hidden"

# ====================================================================== Simulations: column layout
s = wb.create_sheet("Simulations")
cols = []  # (key, header, width, band_key, band_title, kind, comment)


def add(key, head, width, band, band_title, kind="in", comment=None):
    cols.append((key, head, width, band, band_title, kind, comment))


add("id", "Sim ID", 7, "id", "Identification", comment="Unique short ID, e.g. S12")
add("name", "File name", 24, "id", "Identification", comment="PAMsim 'File name': files become <name>_1.sim, ...")
add("purpose", "Purpose / question", 34, "id", "Identification")
add("folder", "Output folder", 28, "id", "Identification")
add("status", "Status", 10, "id", "Identification")
add("prio", "Priority", 7, "id", "Identification", comment="1 = highest")
for k, hd, w, com in (("bx", "Box x [nm]", 8, None), ("by", "Box y [nm]", 8, None), ("bz", "Box z [nm]", 8, "0 = 2D (membrane)"),
                      ("freq", "Sampling freq. [kHz]", 9, None), ("dur", "Duration [s]", 8, None), ("frames", "Frames", 7, None),
                      ("mir", "Microtime range [ns]", 9, None), ("scan", "Scan type", 10, None),
                      ("ncol", "Colours", 7, "Number of colours / PIE lasers (1-3 in this planner)"),
                      ("bg", "Background?", 9, None), ("bgk", "Background [kHz/ch]", 9, None), ("irf", "IRF?", 6, None),
                      ("irfw", "IRF width [ps]", 7, None)):
    add(k, hd, w, "gen", "Measurement", comment=com)
for ch in (1, 2, 3):
    add(f"wr{ch}", f"wr {ch} [nm]", 7, "foc", "Focus per colour",
        comment="Lateral 1/e² radius of the focus of colour 1 (laser 1)" if ch == 1 else None)
for ch in (1, 2, 3):
    add(f"wz{ch}", f"wz {ch} [nm]", 7, "foc", "Focus per colour",
        comment="Axial 1/e² radius of the focus of colour 1 (laser 1)" if ch == 1 else None)
for ax in ("dX", "dY", "dZ"):
    for ch in (1, 2, 3):
        add(f"{ax}{ch}", f"{ax} {ch} [nm]", 7, "foc", "Focus per colour",
            comment=("Shift of the focus of colour 1 from the box centre (0 = perfectly overlapping foci)"
                     if (ax, ch) == ("dX", 1) else None))
add("r012", "R0 1→2", 7, "fret", "Förster radii R0", comment="Förster radius donor colour 1 → acceptor colour 2 (Å or nm)")
add("r013", "R0 1→3", 7, "fret", "Förster radii R0")
add("r023", "R0 2→3", 7, "fret", "Förster radii R0")
SP_FIELDS = [  # key, header, width, kind
    ("nm", "Name", 14, "in"), ("D", "D [µm²/s]", 8, "in"), ("N", "Molecules N", 9, "in"),
    ("b1", "Bright. 1 [kHz]", 8, "in"), ("b2", "Bright. 2 [kHz]", 8, "in"), ("b3", "Bright. 3 [kHz]", 8, "in"),
    ("t1", "Lifetime 1 [ns]", 8, "in"), ("t2", "Lifetime 2 [ns]", 8, "in"), ("t3", "Lifetime 3 [ns]", 8, "in"),
    ("fret", "FRET", 11, "in"), ("R12", "R 1→2", 7, "in"), ("R13", "R 1→3", 7, "in"), ("R23", "R 2→3", 7, "in"),
    ("E12", "E 1→2", 7, "calc"), ("E13", "E 1→3", 7, "calc"), ("E23", "E 2→3", 7, "calc"),
    ("c", "Conc. [nM]", 8, "calc"), ("chk", "Check", 25, "calc"),
]
for sidx in range(1, MAX_SPECIES + 1):
    for k, hd, w, kind in SP_FIELDS:
        add(f"s{sidx}_{k}", hd, w, f"sp{sidx}", f"Species {sidx}", kind,
            "Brightness 0 = no dye of this colour" if k == "b1" and sidx == 1 else None)
for k, hd, w, com in (("nsp", "Species used", 8, None), ("ntot", "Total molecules", 9, None),
                      ("r1", "Exp. rate ch1 [kHz]", 9, None), ("r2", "Exp. rate ch2 [kHz]", 9, None),
                      ("r3", "Exp. rate ch3 [kHz]", 9, None), ("ph", "Photons per file [M]", 9, None),
                      ("mem", "Memory per file [GB]", 9, None),
                      ("xs", "Crosstalk / dir. exc.", 11, f"Switched on for this simulation on the sheet '{XS_NAME}'"),
                      ("check", "Check", 40, None)):
    add(k, hd, w, "chk", "Totals and checks (automatic)", "calc", com)
add("nfiles", "Number of files", 8, "out", "Output")
add("seed", "Seed", 8, "out", "Output", comment="Empty = random")
add("fmt", "Format", 7, "out", "Output")
for k, hd, w in (("date", "Date run", 10), ("runtime", "Run time [min]", 8), ("fitD", "Fitted D [µm²/s]", 9),
                 ("fitc", "Fitted c [nM]", 9), ("notes", "Result notes", 36)):
    add(k, hd, w, "res", "Results")

C = {c[0]: L(i + 1) for i, c in enumerate(cols)}



def draw_header(ws, cols):
    """Band row (1, merged with row 2) and header row (3) of a planner sheet."""
    i = 0
    while i < len(cols):
        band, title = cols[i][3], cols[i][4]
        j = i
        while j + 1 < len(cols) and cols[j + 1][3] == band:
            j += 1
        color = SP_BAND[(int(band[2:]) - 1) % len(SP_BAND)] if band.startswith("sp") else BAND[band]
        ws.merge_cells(start_row=1, start_column=i + 1, end_row=2, end_column=j + 1)
        c = ws.cell(row=1, column=i + 1, value=title)
        c.font = Font(name=F, size=11, bold=True, color="FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center")
        for k in range(i + 1, j + 2):
            for rr in (1, 2):
                ws.cell(row=rr, column=k).fill = PatternFill("solid", fgColor=color)
        if band.startswith("sp"):  # outline group so species blocks can be folded
            ws.column_dimensions.group(L(i + 1), L(j + 1), outline_level=1, hidden=False)
        i = j + 1
    for idx, (k, hd, w, band, title, kind, com) in enumerate(cols, start=1):
        c = ws.cell(row=3, column=idx, value=hd)
        c.fill = HDR; c.font = WHITE_B; c.border = BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if com:
            c.comment = Comment(com, "PAMsim planner")
        ws.column_dimensions[L(idx)].width = w
    ws.row_dimensions[3].height = 40


draw_header(s, cols)
s.freeze_panes = f"{C['purpose']}{FIRST}"
s.sheet_properties.outlinePr.summaryRight = False

# ====================================================================== crosstalk / direct excitation sheet
xs = wb.create_sheet(XS_NAME)
xcols = []
XPAIRS = (("12", 2), ("13", 3), ("23", 3))  # pair, colours needed
xcols.append(("id", "Sim ID", 8, "xid", "Simulation", "in",
              "Pick the Sim ID of a simulation on the Simulations sheet (drop-down). One row per simulation. "
              "Simulations not listed here have no crosstalk and no direct excitation."))
xcols.append(("name", "File name", 30, "xid", "Simulation", "calc", "Looked up from the Simulations sheet"))
xcols.append(("ncol", "Colours", 8, "xid", "Simulation", "calc", "Looked up from the Simulations sheet"))
xcols.append(("ct", "Crosstalk?", 10, "xct", "Crosstalk (dye → channel)", "in",
              "Yes = simulate the crosstalk of this row. No / empty = no crosstalk (the values are kept but ignored)."))
for pair, _ in XPAIRS:
    xcols.append((f"ct{pair}", f"Crosstalk {pair[0]}→{pair[1]}", 10, "xct", "Crosstalk (dye → channel)", "in",
                  f"Photons of dye {pair[0]} detected in channel {pair[1]}, as a fraction of the photons of dye {pair[0]} "
                  f"in its own channel {pair[0]} (e.g. 0.05 = 5 %). Write a fraction, not a percentage."
                  if pair == "12" else None))
xcols.append(("de", "Direct exc.?", 10, "xde", "Direct excitation (laser → dye)", "in",
              "Yes = simulate the direct excitation of this row. No / empty = no direct excitation."))
for pair, _ in XPAIRS:
    xcols.append((f"de{pair}", f"Direct exc. {pair[0]}→{pair[1]}", 10, "xde", "Direct excitation (laser → dye)", "in",
                  f"Excitation of dye {pair[1]} by laser {pair[0]}, as a fraction of its excitation by its own laser "
                  f"{pair[1]} (e.g. 0.05: under laser {pair[0]} dye {pair[1]} is 5 % as bright as under laser {pair[1]}). "
                  f"Only species that carry dye {pair[1]} (brightness > 0) are affected." if pair == "12" else None))
xcols.append(("check", "Check", 46, "chk", "Check (automatic)", "calc", None))
BAND.update({"xid": "2A78D6", "xct": "1B6F8A", "xde": "8A5A1B"})
XC = {c[0]: L(i + 1) for i, c in enumerate(xcols)}
draw_header(xs, xcols)
xs.freeze_panes = f"B{FIRST}"
XR = lambda k: f"{XS}${XC[k]}${FIRST}:${XC[k]}${LAST}"  # noqa: E731  absolute column range on the crosstalk sheet
SR = lambda k: f"Simulations!${C[k]}${FIRST}:${C[k]}${LAST}"  # noqa: E731  same on the Simulations sheet

# ====================================================================== hidden calculation sheet
calc = wb.create_sheet("Calc")
calc["A1"] = "Per-species helper values for the Simulations sheet (do not edit)"
CC = {}
cc = 2
for sidx in range(1, MAX_SPECIES + 1):
    for k in ("rate1", "rate2", "rate3", "maxb", "step", "used"):
        CC[(sidx, k)] = L(cc)
        calc.cell(row=3, column=cc, value=f"S{sidx} {k}")
        cc += 1
# per simulation: crosstalk / direct excitation in effect, focus volumes, rates without and with them
for k in ("xrow", "cton", "deon", "ct12", "ct13", "ct23", "de12", "de13", "de23",
          "V1", "V2", "V3", "R1", "R2", "R3", "E2", "E3"):
    CC[k] = L(cc)
    calc.cell(row=3, column=cc, value=k)
    cc += 1

# ====================================================================== formulas
DEFAULTS = {"bx": 5000, "by": 5000, "bz": 5000, "freq": 1000, "dur": 10, "ncol": 1,
            "wr1": 200, "wr2": 200, "wr3": 200, "wz1": 1000, "wz2": 1000, "wz3": 1000,
            "r012": 63, "r013": 51, "r023": 68}
for r in range(FIRST, LAST + 1):
    for idx, (k, hd, w, band, title, kind, com) in enumerate(cols, start=1):
        c = s.cell(row=r, column=idx)
        c.border = BOX
        c.fill = INPUT if kind == "in" else CALC
        c.font = BLUE if kind == "in" else BLACK
        c.alignment = Alignment(vertical="top", wrap_text=k in ("purpose", "check", "notes", "folder"))
    # empty input cells get the same defaults as PAMsim uses (no #DIV/0! for empty cells)
    def _d(ref, k):
        return f'IF({ref}="",{DEFAULTS[k]},{ref})' if k in DEFAULTS else ref
    S = lambda k: _d(f"Simulations!${C[k]}{r}", k)  # noqa: E731
    g = lambda k: _d(f"${C[k]}{r}", k)  # noqa: E731
    ida = g("id")
    vol = f"({S('bx')}*{S('by')}*IF({S('bz')}>0,{S('bz')},1))"
    for sidx in range(1, MAX_SPECIES + 1):
        sk = lambda k: f"s{sidx}_{k}"  # noqa: E731
        Nn, Dd = S(sk("N")), S(sk("D"))
        used = f"AND({S('id')}<>\"\",OR({S(sk('nm'))}<>\"\",N({Nn})>0))"
        calc[f"{CC[(sidx, 'used')]}{r}"] = f"=IF({used},1,0)"
        for ch in (1, 2, 3):
            B = S(sk(f"b{ch}"))
            wr, wz = S(f"wr{ch}"), S(f"wz{ch}")  # focus of this colour
            prof = f"(PI()/2*{wr}^2*IF({S('bz')}>0,SQRT(PI()/2)*{wz},1))"
            u0 = f"LN({B}/{S('freq')})"
            sat3 = (f"2*PI()*{wr}^2*{wz}/2^1.5*(2/3*{u0}^1.5+SQRT({u0})"
                    f"+{B}/{S('freq')}*SQRT(PI())/2*ERFC(SQRT({u0})))")
            sat2 = f"PI()*{wr}^2/2*({u0}+1)"
            calc[f"{CC[(sidx, f'rate{ch}')]}{r}"] = (
                f"=IF(OR(NOT({used}),N({B})<=0,N({Nn})<=0,{ch}>N({S('ncol')})),0,"
                f"IF({B}<{S('freq')},{Nn}/{vol}*{B}*{prof},{Nn}/{vol}*{S('freq')}*IF({S('bz')}>0,{sat3},{sat2})))")
        calc[f"{CC[(sidx, 'maxb')]}{r}"] = (f"=IF(NOT({used}),0,MAX(N({S(sk('b1'))}),N({S(sk('b2'))}),N({S(sk('b3'))}))"
                                           f"/{S('freq')})")
        calc[f"{CC[(sidx, 'step')]}{r}"] = (f"=IF(OR(NOT({used}),N({Dd})<=0),0,"
                                           f"SQRT(2*{Dd}*1000000/({S('freq')}*1000))/MIN({S('wr1')},{S('wr2')},{S('wr3')}))")
        # FRET efficiencies (visible)
        for pair, r0 in (("12", "r012"), ("13", "r013"), ("23", "r023")):
            R = g(sk(f"R{pair}"))
            s[f"{C[sk(f'E{pair}')]}{r}"] = (f'=IF(OR({g(sk("fret"))}="",{g(sk("fret"))}="No FRET",N({R})<=0,N({g(r0)})<=0),"",'
                                            f'1/(1+({R}/{g(r0)})^6))')
            s[f"{C[sk(f'E{pair}')]}{r}"].number_format = "0.00"
        # concentration
        s[f"{C[sk('c')]}{r}"] = (f'=IF(OR({ida}="",N({g(sk("N"))})<=0),"",IF({g("bz")}>0,'
                                 f'{g(sk("N"))}/({g("bx")}*{g("by")}*{g("bz")})/1E-24/6.022E+23/1E-9,""))')
        s[f"{C[sk('c')]}{r}"].number_format = "0.000"
        mx = f"Calc!{CC[(sidx, 'maxb')]}{r}"
        stp = f"Calc!{CC[(sidx, 'step')]}{r}"
        fr = g(sk("fret"))
        s[f"{C[sk('chk')]}{r}"] = (
            f'=IF(Calc!{CC[(sidx, "used")]}{r}=0,"",'
            f'IF({mx}>=1,"✖ saturated (B ≥ freq.)",'
            f'IF({mx}>0.1,"⚠ B > 10% of freq.",'
            f'IF({mx}=0,"⚠ no brightness",'
            f'IF({stp}>0.1,"⚠ step > 10% of wr",'
            f'IF(OR(AND(N({g(sk("b3"))})>0,{g("ncol")}<3),AND(N({g(sk("b2"))})>0,{g("ncol")}<2)),'
            f'"⚠ uses a colour above Colours = "&{g("ncol")}&" (ignored)",'
            f'IF(AND({fr}<>"",{fr}<>"No FRET",N({g("ncol")})<2),"⚠ FRET needs ≥ 2 colours",'
            f'IF(AND({fr}<>"",{fr}<>"No FRET",N({g(sk("R12"))})+N({g(sk("R13"))})+N({g(sk("R23"))})=0),"⚠ FRET: set R",'
            f'"✔ OK"))))))))')
    # totals
    used_cells = ",".join(f"Calc!{CC[(i, 'used')]}{r}" for i in range(1, MAX_SPECIES + 1))
    s[f"{C['nsp']}{r}"] = f'=IF({ida}="","",SUM({used_cells}))'
    s[f"{C['ntot']}{r}"] = f'=IF({ida}="","",' + "+".join(f"N({g(f's{i}_N')})" for i in range(1, MAX_SPECIES + 1)) + ")"
    # crosstalk / direct excitation of this simulation (0 when switched off or not listed on the crosstalk sheet)
    K = lambda k: f"Calc!${CC[k]}{r}"  # noqa: E731
    calc[f"{CC['xrow']}{r}"] = f'=IF({S("id")}="",0,IFERROR(MATCH({S("id")},{XR("id")},0),0))'
    calc[f"{CC['cton']}{r}"] = f'=IF({K("xrow")}=0,0,IF(INDEX({XR("ct")},{K("xrow")})="Yes",1,0))'
    calc[f"{CC['deon']}{r}"] = f'=IF({K("xrow")}=0,0,IF(INDEX({XR("de")},{K("xrow")})="Yes",1,0))'
    for pair, need in XPAIRS:
        for kind in ("ct", "de"):
            calc[f"{CC[kind + pair]}{r}"] = (f'=IF(AND({K(kind + "on")}=1,{S("ncol")}>={need}),'
                                             f'N(INDEX({XR(kind + pair)},{K("xrow")})),0)')
    for ch in (1, 2, 3):
        calc[f"{CC[f'V{ch}']}{r}"] = f"=PI()/2*{S(f'wr{ch}')}^2*IF({S('bz')}>0,SQRT(PI()/2)*{S(f'wz{ch}')},1)"
        calc[f"{CC[f'R{ch}']}{r}"] = "=SUM(" + ",".join(f"Calc!{CC[(i, f'rate{ch}')]}{r}"
                                                       for i in range(1, MAX_SPECIES + 1)) + ")"
    # dye 2 / dye 3 photons incl. direct excitation by the bluer laser(s) in that laser's focus (unsaturated)
    calc[f"{CC['E2']}{r}"] = f'={K("R2")}*(1+IF({K("V2")}>0,{K("de12")}*{K("V1")}/{K("V2")},0))'
    calc[f"{CC['E3']}{r}"] = (f'={K("R3")}*(1+IF({K("V3")}>0,({K("de13")}*{K("V1")}+{K("de23")}*{K("V2")})'
                              f'/{K("V3")},0))')
    rate_f = {1: K("R1"), 2: f'{K("E2")}+{K("ct12")}*{K("R1")}',
              3: f'{K("E3")}+{K("ct13")}*{K("R1")}+{K("ct23")}*{K("E2")}'}
    for ch in (1, 2, 3):
        s[f"{C[f'r{ch}']}{r}"] = f'=IF({ida}="","",{rate_f[ch]})'
        s[f"{C[f'r{ch}']}{r}"].number_format = '0.0;-0.0;"-"'
    on_ct = f'AND({K("cton")}=1,{S("ncol")}>=2)'
    on_de = f'AND({K("deon")}=1,{S("ncol")}>=2)'
    s[f"{C['xs']}{r}"] = (f'=IF({ida}="","",IF(AND({on_ct},{on_de}),"Crosstalk + dir. exc.",'
                          f'IF({on_ct},"Crosstalk",IF({on_de},"Direct exc.","–"))))')
    s[f"{C['ph']}{r}"] = (f'=IF({ida}="","",({g("r1")}+{g("r2")}+{g("r3")}+IF({g("bg")}="Yes",N({g("bgk")})*N({g("ncol")}),0))'
                          f'*1000*{g("dur")}/1000000)')
    s[f"{C['ph']}{r}"].number_format = "0.00"
    s[f"{C['mem']}{r}"] = f'=IF({ida}="","",{g("ph")}*1000000*12/1000000000)'
    s[f"{C['mem']}{r}"].number_format = "0.000"
    chks = [g(f"s{i}_chk") for i in range(1, MAX_SPECIES + 1)]
    nred = "+".join(f'(LEFT({c},1)="✖")' for c in chks)
    nor = "+".join(f'(LEFT({c},1)="⚠")' for c in chks)
    xmsg = f'IF({K("xrow")}>0,INDEX({XR("check")},{K("xrow")}),"")'  # check of this simulation on the crosstalk sheet
    s[f"{C['check']}{r}"] = (
        f'=IF({ida}="","",IF({g("nsp")}=0,"⚠ no species entered",'
        f'IF(({nred})>0,"✖ "&({nred})&" species saturated - see the species checks",'
        f'IF(MIN({g("bx")},{g("by")})<8*MAX({g("wr1")},{g("wr2")},{g("wr3")}),"⚠ box too small laterally (< 8·wr)",'
        f'IF(AND({g("bz")}>0,{g("bz")}<3*MAX({g("wz1")},{g("wz2")},{g("wz3")})),"⚠ box too small axially (< 3·wz)",'
        f'IF({g("mem")}>8,"⚠ > 8 GB per file: shorten the duration and use several files",'
        f'IF(({nor})>0,"⚠ "&({nor})&" species with a warning",'
        f'IF(OR(LEFT({xmsg},1)="✖",LEFT({xmsg},1)="⚠"),"⚠ crosstalk sheet: "&MID({xmsg},3,200),"✔ OK"))))))))')
calc.sheet_state = "hidden"

# ====================================================================== validation and formatting
def dv(ref, cells):
    v = DataValidation(type="list", formula1=ref, allow_blank=True)
    s.add_data_validation(v)
    for c in cells:
        v.add(c)


rng = lambda k: f"{C[k]}{FIRST}:{C[k]}{LAST}"  # noqa: E731
dv("Lists!$A$2:$A$7", [rng("status")])
dv("Lists!$B$2:$B$5", [rng("scan")])
dv("Lists!$C$2:$C$3", [rng("bg"), rng("irf")])
dv("Lists!$E$2:$E$4", [rng("fmt")])
dv("Lists!$D$2:$D$4", [rng(f"s{i}_fret") for i in range(1, MAX_SPECIES + 1)])
vcol = DataValidation(type="whole", operator="between", formula1="1", formula2="3", allow_blank=True)
vcol.error = "This planner handles 1-3 colours"; s.add_data_validation(vcol); vcol.add(rng("ncol"))
vpos = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
vpos.error = "Must be >= 0"; s.add_data_validation(vpos)
for i in range(1, MAX_SPECIES + 1):
    for k in ("D", "N", "b1", "b2", "b3", "t1", "t2", "t3", "R12", "R13", "R23"):
        vpos.add(rng(f"s{i}_{k}"))


def traffic(col_range, first_cell, ws=s):
    ws.conditional_formatting.add(col_range, FormulaRule(formula=[f'LEFT({first_cell},1)="✖"'], fill=RED,
                                                         font=Font(name=F, color="9C1C1C", bold=True)))
    ws.conditional_formatting.add(col_range, FormulaRule(formula=[f'LEFT({first_cell},1)="⚠"'], fill=ORANGE,
                                                         font=Font(name=F, color="8A4B00")))
    ws.conditional_formatting.add(col_range, FormulaRule(formula=[f'LEFT({first_cell},1)="✔"'], fill=GREEN,
                                                         font=Font(name=F, color="1E6B1E")))


traffic(rng("check"), f"${C['check']}{FIRST}")
for i in range(1, MAX_SPECIES + 1):
    traffic(rng(f"s{i}_chk"), f"${C[f's{i}_chk']}{FIRST}")
st = C["status"]
for val, colr in (("Done", "D5EFD5"), ("Analysed", "BFE3C6"), ("Running", "D6E6FA"), ("Discarded", "E3E3E3")):
    s.conditional_formatting.add(rng("status"), FormulaRule(formula=[f'${st}{FIRST}="{val}"'],
                                                            fill=PatternFill("solid", fgColor=colr)))
s.auto_filter.ref = f"A3:{L(len(cols))}{LAST}"

# ====================================================================== crosstalk sheet: formulas and formatting
for r in range(FIRST, LAST + 1):
    for idx, (k, hd, w, band, title, kind, com) in enumerate(xcols, start=1):
        c = xs.cell(row=r, column=idx)
        c.border = BOX
        c.fill = INPUT if kind == "in" else CALC
        c.font = BLUE if kind == "in" else BLACK
        c.alignment = Alignment(vertical="top")
    x = lambda k: f"${XC[k]}{r}"  # noqa: E731
    found = f'MATCH({x("id")},{SR("id")},0)'
    xs[x("name").replace("$", "")] = f'=IF({x("id")}="","",IFERROR(INDEX({SR("name")},{found})&"",""))'
    xs[x("ncol").replace("$", "")] = f'=IF({x("id")}="","",IFERROR(MAX(1,N(INDEX({SR("ncol")},{found}))),""))'
    on_ct, on_de = f'{x("ct")}="Yes"', f'{x("de")}="Yes"'
    vals = lambda kind, pairs: "+".join(f'N({x(kind + p)})' for p in pairs)  # noqa: E731
    big = ",".join(f'N({x(kind + p)})' for kind in ("ct", "de") for p, _ in XPAIRS)
    xs[x("check").replace("$", "")] = (
        f'=IF({x("id")}="","",'
        f'IF(COUNTIF({SR("id")},{x("id")})=0,"✖ Sim ID not found on the Simulations sheet",'
        f'IF(COUNTIF(${XC["id"]}${FIRST}:${XC["id"]}${LAST},{x("id")})>1,"✖ Sim ID entered twice on this sheet",'
        f'IF(AND(NOT({on_ct}),NOT({on_de})),"– off (set Crosstalk? and/or Direct exc.? to Yes)",'
        f'IF(N({x("ncol")})<2,"⚠ needs ≥ 2 colours - ignored",'
        f'IF(MAX({big})>1,"⚠ value > 1: write fractions (0.05 = 5 %)",'
        f'IF(AND({on_ct},{vals("ct", ("12", "13", "23"))}=0),"⚠ crosstalk switched on but all values are 0",'
        f'IF(AND({on_de},{vals("de", ("12", "13", "23"))}=0),"⚠ direct exc. switched on but all values are 0",'
        f'IF(AND(N({x("ncol")})<3,OR(AND({on_ct},{vals("ct", ("13", "23"))}>0),AND({on_de},{vals("de", ("13", "23"))}>0))),'
        f'"⚠ 1→3 and 2→3 ignored (Colours < 3)",'
        f'"✔ on: "&IF({on_ct},"crosstalk","")&IF(AND({on_ct},{on_de})," + ","")&IF({on_de},"direct excitation","")'
        f')))))))))')
xrng = lambda k: f"{XC[k]}{FIRST}:{XC[k]}{LAST}"  # noqa: E731
v = DataValidation(type="list", formula1=SR("id"), allow_blank=True)
xs.add_data_validation(v); v.add(xrng("id"))
v = DataValidation(type="list", formula1="Lists!$C$2:$C$3", allow_blank=True)
xs.add_data_validation(v); v.add(xrng("ct")); v.add(xrng("de"))
v = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
v.error = "Must be >= 0"; xs.add_data_validation(v)
for pair, _ in XPAIRS:
    v.add(xrng(f"ct{pair}")); v.add(xrng(f"de{pair}"))
    for kind in ("ct", "de"):
        for r in range(FIRST, LAST + 1):
            xs[f"{XC[kind + pair]}{r}"].number_format = "0.000"
traffic(xrng("check"), f"${XC['check']}{FIRST}", ws=xs)
for k in ("ct", "de"):
    xs.conditional_formatting.add(xrng(k), FormulaRule(formula=[f'${XC[k]}{FIRST}="Yes"'], fill=GREEN,
                                                       font=Font(name=F, color="1E6B1E", bold=True)))
xs.auto_filter.ref = f"A3:{L(len(xcols))}{LAST}"

# ====================================================================== pre-filled rows
base = "/Fablab 4/Users/Fabio/PhD/Projects/3C_FCCS/Simulation"
cpx = base + "/Complex"
lt3, cx = [4, 3, 2], [4, 3, 2.5]
sims = [
    ("S00", "example_3C_realistic", "EXAMPLE: realistic 3-colour FCCS mixture incl. a static FRET pair (brightness < 10% of freq.)",
     base, "Idea", 1, 60, [("triple complex", 35, 300, [100, 70, 50], lt3, None),
                           ("blue-green", 35, 200, [100, 70, 0], lt3, None),
                           ("green-red", 35, 200, [0, 70, 50], lt3, None),
                           ("FRET pair 1-2", 50, 100, [100, 70, 0], lt3, ("Static FRET", 50, None, None))]),
    ("S01", "Simulation_800N_15D_200wr_1000wz_60s_60ns", "Just to see if the code works", base, "Planned", 2, 60,
     [("Species 1", 15, 800, [50, 50, 50], lt3, None)]),
    ("S02", "Simulation_100Nx3_200wr_1000wz_60s_60ns_3colorXspecies", "To verify that no 3C cross-correlation would arise",
     base, "Planned", 2, 60, [("Species 1", 50, 800, [50, 50, 0], lt3, None), ("Species 2", 50, 100, [50, 0, 50], lt3, None),
                              ("Species 3", 50, 100, [0, 50, 50], lt3, None)]),
    ("S03", "Simulation_complexes_1st_10N_200wr_1000wz_10s_60ns_1000kHz", "To test the code in a crowded environment", cpx,
     "Planned", 3, 10, [("Species 1", 400, 10, [1e4, 0, 0], cx, None), ("Species 2", 400, 10, [0, 1e4, 0], cx, None),
                        ("Species 3", 400, 10, [0, 0, 1e4], cx, None), ("Species 4", 33, 10, [1e4, 1e4, 1e4], cx, None),
                        ("Species 5", 35, 10, [1e4, 1e4, 0], cx, None), ("Species 6", 33, 10, [2e4, 1e4, 0], cx, None)]),
    ("S04", "Simulation_complexes_2nd_10N_200wr_1000wz_10s_60ns_1000kHz", "To test the code in a crowded environment", cpx,
     "Planned", 3, 10, [("Species 1", 400, 10, [1e4, 0, 0], cx, None), ("Species 2", 400, 10, [0, 7e3, 0], cx, None),
                        ("Species 3", 400, 10, [0, 0, 5e3], cx, None), ("Species 4", 33, 10, [1e4, 7e3, 5e3], cx, None),
                        ("Species 5", 33, 10, [2e4, 7e3, 0], cx, None)]),
    ("S05", "Simulation_complexes_3rd_10N_200wr_1000wz_10s_60ns_1000kHz",
     "To test whether the code would give rise to a 3C cross-correlation when wrong", cpx, "Planned", 3, 10,
     [("Species 1", 400, 10, [1e4, 0, 0], cx, None), ("Species 2", 400, 10, [0, 7e3, 0], cx, None),
      ("Species 3", 400, 10, [0, 0, 5e3], cx, None), ("Species 4", 33, 10, [2e4, 7e3, 0], cx, None),
      ("Species 5", 33, 10, [1e4, 7e3, 5e3], cx, None)]),
    ("S06", "Simulation_complexes_4th_10N_200wr_1000wz_10s_60ns_1000kHz",
     "To test whether the code would give rise to a 3C cross-correlation when wrong", cpx, "Planned", 3, 10,
     [("Species 1", 400, 10, [1e4, 0, 0], cx, None), ("Species 2", 400, 10, [0, 7e3, 0], cx, None),
      ("Species 3", 400, 10, [0, 0, 5e3], cx, None), ("Species 4", 33, 10, [2e4, 7e3, 0], cx, None),
      ("Species 5", 33, 10, [1e4, 0, 5e3], cx, None), ("Species 6", 33, 10, [0, 7e3, 1e4], cx, None)]),
    ("S07", "Crazy", "What happens to 3C if the environment is very crowded", cpx, "Planned", 3, 30,
     [("Species 1", 35, 100, [1e3, 700, 500], lt3, None), ("Species 2", 35, 1000, [1e3, 0, 500], lt3, None),
      ("Species 3", 35, 500, [1e3, 700, 0], lt3, None)]),
    ("S08", "Crazy_2", "you need to shorten the lag time", cpx, "Planned", 3, 30,
     [("Species 1", 35, 100, [1e3, 700, 500], lt3, None), ("Species 2", 35, 1000, [1e3, 0, 500], lt3, None),
      ("Species 3", 35, 2000, [2e3, 700, 0], lt3, None)]),
    ("S09", "Crazy_3", "", cpx, "Planned", 3, 30,
     [("Species 1", 35, 10, [1e3, 700, 500], lt3, None), ("Species 2", 35, 10, [1e3, 0, 500], lt3, None),
      ("Species 3", 35, 10, [2e3, 700, 0], lt3, None)]),
    ("S10", "Crazy_4", "", cpx, "Planned", 3, 5,
     [("Species 1", 35, 1000, [1e3, 700, 500], lt3, None), ("Species 2", 35, 1000, [1e3, 0, 500], lt3, None),
      ("Species 3", 35, 1000, [2e3, 700, 0], lt3, None)]),
    ("S11", "Crazy_2_checkup_final", "longer simulation with same concept as crazy_2 but better statistics", cpx + "/Crazy_2",
     "Planned", 3, 150, [("Species 1", 35, 200, [1e4, 7e3, 5e3], lt3, None), ("Species 2", 35, 2000, [1e4, 0, 5e3], lt3, None),
                         ("Species 3", 35, 4000, [2e4, 7e3, 0], lt3, None)]),
]
if ARGS.empty or OLD:
    sims = []
for k, (sid, name, purpose, folder, status, prio, dur, species) in enumerate(sims):
    r = FIRST + k
    vals = {"id": sid, "name": name, "purpose": purpose, "folder": folder, "status": status, "prio": prio,
            "bx": 5000, "by": 5000, "bz": 5000, "freq": 1000, "dur": dur, "frames": 1, "mir": 60, "scan": "Point",
            "ncol": 3, "bg": "No", "bgk": 0.2, "irf": "No", "irfw": 250,
            "wr1": 200, "wr2": 200, "wr3": 200, "wz1": 1000, "wz2": 1000, "wz3": 1000,
            "dX1": 0, "dX2": 0, "dX3": 0, "dY1": 0, "dY2": 0, "dY3": 0, "dZ1": 0, "dZ2": 0, "dZ3": 0,
            "r012": 63, "r013": 51, "r023": 68, "nfiles": 1, "fmt": ".sim",
            "notes": "Example row - shows the expected format" if sid == "S00"
            else "from Simulation_parameters.xlsx (box size assumed)"}
    if sid == "S00":
        vals.update({"wr1": 180, "wr2": 210, "wr3": 240, "wz1": 900, "wz2": 1050, "wz3": 1200, "dZ3": 100})
    for key, v in vals.items():
        s[f"{C[key]}{r}"] = v
    for sidx, (nm, D, N, b, lt, fret) in enumerate(species, start=1):
        if sidx > MAX_SPECIES:
            break
        sk = lambda kk: C[f"s{sidx}_{kk}"]  # noqa: E731
        s[f"{sk('nm')}{r}"] = nm; s[f"{sk('D')}{r}"] = D; s[f"{sk('N')}{r}"] = N
        for ch in range(3):
            s[f"{sk(f'b{ch + 1}')}{r}"] = b[ch]
            s[f"{sk(f't{ch + 1}')}{r}"] = lt[ch]
        if fret:
            s[f"{sk('fret')}{r}"] = fret[0]
            for pair, val in zip(("R12", "R13", "R23"), fret[1:]):
                if val is not None:
                    s[f"{sk(pair)}{r}"] = val
        else:
            s[f"{sk('fret')}{r}"] = "No FRET"



def copy_old(ws, layout, records):
    """Write the input cells of an old planner into ``ws``; species blocks by (band, header), others by header."""
    renamed = {"wr [nm]": ("wr 1 [nm]", "wr 2 [nm]", "wr 3 [nm]"),   # planners with one focus for all colours
               "wz [nm]": ("wz 1 [nm]", "wz 2 [nm]", "wz 3 [nm]")}
    for k, rec in enumerate(records):
        by_head = {}
        for (band, head), v in rec.items():
            if not band.startswith("Species "):
                for h in renamed.get(head, (head,)):
                    by_head.setdefault(h, v)
        for idx, (key, head, w, band, title, kind, com) in enumerate(layout, start=1):
            if kind != "in":
                continue
            v = rec.get((title, head)) if title.startswith("Species ") else by_head.get(head)
            if v is not None:
                ws.cell(row=FIRST + k, column=idx, value=v)


if OLD:
    copy_old(s, cols, OLD["sim"])
    copy_old(xs, xcols, OLD["xs"])

# example on the crosstalk sheet: realistic values for S00, switched off (set the switches to Yes to use them)
if not ARGS.empty and not OLD:
    for key, val in {"id": "S00", "ct": "No", "ct12": 0.08, "ct13": 0.01, "ct23": 0.05,
                     "de": "No", "de12": 0.05, "de13": 0.01, "de23": 0.06}.items():
        xs[f"{XC[key]}{FIRST}"] = val

for name in ("How to use", "Column guide", "Simulations", XS_NAME, "Lists", "Calc"):
    wb.move_sheet(name, offset=len(wb.sheetnames))  # to the end, in this order
wb.active = wb.sheetnames.index("Simulations")
wb.save(OUT)
print(f"saved {OUT}  ({N_ROWS} simulation rows, {MAX_SPECIES} species blocks, {len(cols)} columns)")
if OLD:
    print(f"copied {len(OLD['sim'])} simulation(s) and {len(OLD['xs'])} crosstalk row(s) from {ARGS.source}")
print("Open it in Excel; the calculated cells fill in when Excel opens the file.")
